# -*- coding: utf-8 -*-
"""第二类账号：发票明细查看（只读）+ 明细导出。"""

from __future__ import annotations

from datetime import datetime as _dt
from io import BytesIO

from fastapi import APIRouter, Depends, Request
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import or_
from sqlalchemy.orm import Session, selectinload

from ..auth import require_type2
from ..models import (
    CHECK_FAILED,
    CHECK_PASSED,
    CHECK_PENDING,
    Invoice,
    InvoiceItem,
    User,
    get_db,
)
from ..service import (
    attach_detail_extras,
    export_bill_no,
    export_container,
    export_map_for,
    get_exporter,
)
from ..templates_setup import templates

router = APIRouter(tags=["details"])


def _missing_prices_for_invoice(invoice: Invoice) -> list[str]:
    """该发票明细中价格表查不到单价的面料类型。"""
    from core.price_table import lookup_price

    from ..service import company_for_price

    prices = get_exporter().prices
    company = company_for_price(invoice)
    found: set[str] = set()
    for item in invoice.items:
        ft = (item.fabric_type or "").strip()
        if ft and ft not in found and lookup_price(prices, company, ft) == 0.0:
            found.add(ft)
    return sorted(found)


@router.get("/details")
async def detail_list(
    request: Request,
    invoice_id: int = 0,
    q: str = "",
    inv_q: str = "",
    check_status: str = "",
    page: int = 1,
    user: User = Depends(require_type2),
    db: Session = Depends(get_db),
) -> object:
    inv_query = db.query(Invoice)
    if check_status in (CHECK_PENDING, CHECK_PASSED, CHECK_FAILED):
        inv_query = inv_query.filter(Invoice.check_status == check_status)
    if inv_q.strip():
        like = f"%{inv_q.strip()}%"
        inv_query = inv_query.filter(
            or_(
                Invoice.invoice_no.ilike(like),
                Invoice.delivery_to.ilike(like),
                Invoice.bill_to.ilike(like),
            )
        )
    # 下拉只取必要字段，避免 selectin 连带加载每张发票的全部明细行
    invoices = (
        inv_query.with_entities(
            Invoice.id, Invoice.invoice_no, Invoice.delivery_to, Invoice.transport
        )
        .order_by(Invoice.uploaded_at.desc(), Invoice.id.desc())
        .limit(500)
        .all()
    )

    # 搜索发票号唯一命中时自动选中，免去再点一次下拉
    if inv_q.strip() and not invoice_id and len(invoices) == 1:
        invoice_id = invoices[0].id

    # 是否处于"按发票筛选"状态：有搜索词或有核对状态筛选
    search_active = bool(inv_q.strip()) or check_status in (
        CHECK_PENDING,
        CHECK_PASSED,
        CHECK_FAILED,
    )
    matched_ids = [i.id for i in invoices]

    query = db.query(InvoiceItem)
    if invoice_id:
        query = query.filter(InvoiceItem.invoice_id == invoice_id)
    elif search_active:
        # 有搜索/筛选时，明细表只显示匹配发票的行（多张匹配则合并显示）
        if not matched_ids:
            query = query.filter(InvoiceItem.invoice_id.is_(None))
        else:
            query = query.filter(InvoiceItem.invoice_id.in_(matched_ids))
    if q.strip():
        like = f"%{q.strip()}%"
        query = query.filter(
            or_(
                InvoiceItem.description.ilike(like),
                InvoiceItem.fabric_type.ilike(like),
                InvoiceItem.po_no.ilike(like),
                InvoiceItem.colour.ilike(like),
            )
        )

    total = query.count()
    page = max(page, 1)
    size = 50
    rows = (
        query.options(selectinload(InvoiceItem.invoice))
        .order_by(InvoiceItem.invoice_id.desc(), InvoiceItem.row.asc())
        .offset((page - 1) * size)
        .limit(size)
        .all()
    )
    pages = max((total + size - 1) // size, 1)

    # 选中发票单独加载完整对象（含明细，供抬头与单价缺失提示）
    current = db.get(Invoice, invoice_id) if invoice_id else None

    attach_detail_extras(rows)
    current_bill_no = (export_bill_no(current) or "-") if current else "-"
    current_container = (export_container(current) or "-") if current else "-"
    current_missing_prices = _missing_prices_for_invoice(current) if current else []

    status_ids = list(matched_ids)
    if invoice_id:
        status_ids.append(invoice_id)
    export_map = export_map_for(db, status_ids)
    generated_ids = set(export_map)
    current_export = export_map.get(invoice_id) if invoice_id else None

    return templates.TemplateResponse(
        request,
        "details.html",
        {
            "user": user,
            "invoices": invoices,
            "current": current,
            "rows": rows,
            "current_bill_no": current_bill_no,
            "current_container": current_container,
            "current_export": current_export,
            "generated_ids": generated_ids,
            "current_missing_prices": current_missing_prices,
            "invoice_id": invoice_id,
            "invoice_match_count": len(invoices),
            "search_active": search_active,
            "q": q,
            "inv_q": inv_q,
            "check_status": check_status,
            "page": page,
            "pages": pages,
            "total": total,
            "flash": request.session.pop("flash", None),
        },
    )


@router.get("/details/export")
async def details_export(
    invoice_id: int = 0,
    q: str = "",
    inv_q: str = "",
    check_status: str = "",
    user: User = Depends(require_type2),
    db: Session = Depends(get_db),
) -> object:
    """导出发票明细为 Excel（无模板，按数据结构生成，与明细页口径一致）。"""
    inv_query = db.query(Invoice)
    if check_status in (CHECK_PENDING, CHECK_PASSED, CHECK_FAILED):
        inv_query = inv_query.filter(Invoice.check_status == check_status)
    if inv_q.strip():
        like = f"%{inv_q.strip()}%"
        inv_query = inv_query.filter(
            or_(
                Invoice.invoice_no.ilike(like),
                Invoice.delivery_to.ilike(like),
                Invoice.bill_to.ilike(like),
            )
        )
    invoices = (
        inv_query.with_entities(Invoice.id)
        .order_by(Invoice.uploaded_at.desc(), Invoice.id.desc())
        .limit(500)
        .all()
    )
    matched_ids = [i.id for i in invoices]
    search_active = bool(inv_q.strip()) or check_status in (
        CHECK_PENDING,
        CHECK_PASSED,
        CHECK_FAILED,
    )

    query = db.query(InvoiceItem)
    if invoice_id:
        query = query.filter(InvoiceItem.invoice_id == invoice_id)
    elif search_active:
        if not matched_ids:
            query = query.filter(InvoiceItem.invoice_id.is_(None))
        else:
            query = query.filter(InvoiceItem.invoice_id.in_(matched_ids))
    if q.strip():
        like = f"%{q.strip()}%"
        query = query.filter(
            or_(
                InvoiceItem.description.ilike(like),
                InvoiceItem.fabric_type.ilike(like),
                InvoiceItem.po_no.ilike(like),
                InvoiceItem.colour.ilike(like),
            )
        )
    rows = (
        query.options(selectinload(InvoiceItem.invoice))
        .order_by(InvoiceItem.invoice_id.desc(), InvoiceItem.row.asc())
        .limit(20000)
        .all()
    )
    attach_detail_extras(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "发票明细"
    headers = [
        "Fabric Invoice No.", "柜号", "Invoice No.", "PO No.", "Description", "Colour",
        "Size", "Fabric Type", "CARTON", "净重", "毛重", "Length",
        "Quantity", "Unit Price", "Amount", "报关净重", "报关单价",
        "报关金额", "备注",
    ]
    header_font = Font(bold=True)
    orange_fill = PatternFill(fill_type="solid", start_color="F8CBAD")
    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for it in rows:
        inv = it.invoice
        remark_parts = []
        if not it.in_export:
            remark_parts.append("* 不在导出范围(17-44行)")
        if it.up == 0:
            remark_parts.append("无单价（需在表格 I 列手动填写）")
        ws.append([
            inv.invoice_no if inv else None,
            it.container_no,
            it.export_bill_no,
            it.po_no,
            it.description,
            it.colour,
            it.size,
            it.fabric_type,
            it.carton,
            it.net_weight,
            it.gross_weight,
            it.length,
            it.quantity,
            it.unit_price,
            it.amount,
            it.cw,
            it.up,
            it.ca,
            "；".join(remark_parts),
        ])
        if it.up == 0:
            ws.cell(row=ws.max_row, column=17).fill = orange_fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [14, 16, 14, 14, 40, 12, 12, 14, 10, 10, 10, 10, 10, 11, 11, 11, 11, 11, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    buffer = BytesIO()
    wb.save(buffer)
    wb.close()
    name = f"invoice_details_{_dt.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )


