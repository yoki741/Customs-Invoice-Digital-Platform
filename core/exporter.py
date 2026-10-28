# -*- coding: utf-8 -*-
"""发票 / 装箱单生成器。

逻辑移植自 invoice_gui.InvoiceProcessor（EXE），保证 Web 导出结果与 EXE 一致：
- 依据运输方式选择模板（海运 / 空运 / 陆运）
- 预填所有已知数据（Fabric Type 汇总、明细、单价、金额、重量、Consignee、REF）
- 留空由人工填写的字段（date / bill_no / ship_mode / 起运地 / 口岸 / 集装箱）
- 通过 zip 层注入 vbaProject.bin 生成含宏的 .xlsm
"""

from __future__ import annotations

import re
import time
import zipfile
from collections import defaultdict
from copy import copy
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Optional

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, PatternFill
from openpyxl.workbook.workbook import Workbook

from .config import TEMPLATE_DIR, TRANSPORT_TEMPLATES
from .price_table import get_transport_type, load_prices, lookup_price, weight_ratio
from .source_reader import SourceData


@dataclass
class ManualFields:
    """由人工填写的字段（对应 EXE 的"手动输入信息"）。"""

    date: str = ""
    bill_no: str = ""
    supplier: str = ""
    ship_mode: str = ""
    place_receipt: str = ""
    port_loading: str = ""
    container_info: str = ""
    flight_no: str = ""
    # 单价手动覆盖：{面料类型(大写): 单价}，缺省时按单价表自动读取
    unit_prices: dict[str, float] = field(default_factory=dict)

    @classmethod
    def from_dict(cls, raw: Optional[dict[str, Any]]) -> "ManualFields":
        raw = raw or {}
        return cls(
            date=str(raw.get("date") or "").strip(),
            bill_no=str(raw.get("bill_no") or "").strip(),
            supplier=str(raw.get("supplier") or "").strip(),
            ship_mode=str(raw.get("ship_mode") or "").strip(),
            place_receipt=str(raw.get("place_receipt") or "").strip(),
            port_loading=str(raw.get("port_loading") or "").strip(),
            container_info=str(raw.get("container_info") or "").strip(),
            flight_no=str(raw.get("flight_no") or "").strip(),
            unit_prices=parse_unit_prices(raw.get("unit_prices")),
        )

    def as_dict(self) -> dict[str, str]:
        return {
            "date": self.date,
            "bill_no": self.bill_no,
            "supplier": self.supplier,
            "ship_mode": self.ship_mode,
            "place_receipt": self.place_receipt,
            "port_loading": self.port_loading,
            "container_info": self.container_info,
            "flight_no": self.flight_no,
        }


def parse_unit_prices(raw: Any) -> dict[str, float]:
    """解析单价覆盖值（支持 JSON 字符串或 dict），返回 {大写面料类型: 单价}。"""
    import json

    data: Any = raw
    if isinstance(raw, str):
        try:
            data = json.loads(raw or "{}")
        except (TypeError, ValueError):
            data = {}
    if not isinstance(data, dict):
        return {}
    prices: dict[str, float] = {}
    for key, value in data.items():
        key_text = str(key or "").strip().upper()
        if not key_text:
            continue
        if value is None or str(value).strip() == "":
            continue
        try:
            prices[key_text] = float(value)
        except (TypeError, ValueError):
            continue
    return prices


@dataclass
class BuildResult:
    """build() 的元信息。"""

    workbook: Workbook
    transport: str = "海运"
    company: str = "Eastex"
    bill_no: str = ""
    file_stem: str = ""
    customs_weight: int = 0
    total_amount: float = 0.0
    container_row: Optional[int] = None
    total_row: int = 37
    ref_row: int = 24
    last_data_row: int = 24
    messages: list[str] = field(default_factory=list)
    ok: bool = True


def _split_label(value: Any) -> tuple[str, str]:
    """把 'Port of loading: NINGBO, CHINA' 拆成 ('Port of loading: ', 'NINGBO, CHINA')。"""
    text = str(value or "")
    if ":" in text:
        idx = text.index(":")
        prefix = text[: idx + 1]
        if len(text) > idx + 1 and text[idx + 1] == " ":
            prefix += " "
        rest = text[idx + 1 :].strip()
        return prefix, rest
    return "Origin: ", text.strip()


def copy_sheet_with_format(src_wb: Workbook, src_name: str, dst_wb: Workbook, dst_name: str) -> None:
    """带格式复制工作表（含合并单元格、列宽、行高）。"""
    src_ws = src_wb[src_name]
    dst_ws = dst_wb[dst_name] if dst_name in dst_wb.sheetnames else dst_wb.create_sheet(dst_name)

    for row in src_ws.iter_rows():
        for cell in row:
            new_cell = dst_ws.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                new_cell.font = copy(cell.font)
                new_cell.border = copy(cell.border)
                new_cell.fill = copy(cell.fill)
                new_cell.number_format = copy(cell.number_format)
                new_cell.alignment = copy(cell.alignment)

    for merged_range in src_ws.merged_cells.ranges:
        dst_ws.merge_cells(str(merged_range))

    for col_letter, dim in src_ws.column_dimensions.items():
        if dim.width:
            dst_ws.column_dimensions[col_letter].width = dim.width

    for row_num, dim in src_ws.row_dimensions.items():
        if dim.height:
            dst_ws.row_dimensions[row_num].height = dim.height


TRACE_SHEET = "_PLATFORM_TRACE"


def add_trace_sheet(workbook: Workbook, meta: dict[str, str]) -> None:
    """往工作簿注入深度隐藏的追踪表（veryHidden，Excel 界面不可见）。

    记录导出任务、导出人、时间、版本等，供文件回传时校验来源。
    """
    if TRACE_SHEET in workbook.sheetnames:
        del workbook[TRACE_SHEET]
    ws = workbook.create_sheet(TRACE_SHEET)
    ws.sheet_state = "veryHidden"
    ws["A1"] = "报关发票数字化平台-文件追踪信息（仅供追溯，请勿修改）"
    row = 3
    for key, value in meta.items():
        ws[f"A{row}"] = str(key)
        ws[f"B{row}"] = str(value)
        row += 1
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 80


def read_trace_sheet(file_path: Path) -> dict[str, str]:
    """读取文件中的追踪表内容；无追踪表或读取失败返回空 dict。"""
    try:
        wb = openpyxl.load_workbook(str(file_path), read_only=True, keep_vba=False)
    except Exception:
        return {}
    try:
        if TRACE_SHEET not in wb.sheetnames:
            return {}
        ws = wb[TRACE_SHEET]
        meta: dict[str, str] = {}
        for row in ws.iter_rows(min_row=2, max_row=40, max_col=2):
            key, value = row[0].value, row[1].value
            if key and value:
                meta[str(key).strip()] = str(value).strip()
        return meta
    finally:
        wb.close()


def inject_vba(xlsx_path: Path, carrier: Path, out_xlsm: Path) -> bool:
    """把源 .xlsm 的 VBA 工程注入 openpyxl 生成的 .xlsx，产出 .xlsm。

    全程在 zip 层操作，无需密码、无需 Excel。
    """
    try:
        with zipfile.ZipFile(carrier) as cz:
            vba_bin = cz.read("xl/vbaProject.bin")

        with zipfile.ZipFile(xlsx_path) as z:
            content_types = z.read("[Content_Types].xml").decode("utf-8")
            rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
            items = [(n, z.read(n)) for n in z.namelist()]

        content_types = content_types.replace(
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.'
            'openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>',
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.'
            'ms-excel.sheet.macroEnabled.main+xml"/>',
        )
        if "vbaProject" not in content_types:
            content_types = content_types.replace(
                "</Types>",
                '<Override PartName="/xl/vbaProject.bin" ContentType='
                '"application/vnd.ms-office.vbaProject"/></Types>',
            )
        ids = [int(m.group(1)) for m in re.finditer(r'Id="rId(\d+)"', rels)]
        next_id = (max(ids) + 1) if ids else 1
        rels = rels.replace(
            "</Relationships>",
            f'<Relationship Id="rId{next_id}" Type="http://schemas.microsoft.com/office/'
            f'2006/relationships/vbaProject" Target="vbaProject.bin"/></Relationships>',
        )

        with zipfile.ZipFile(out_xlsm, "w", zipfile.ZIP_DEFLATED) as oz:
            for name, data in items:
                if name == "[Content_Types].xml":
                    oz.writestr(name, content_types)
                elif name == "xl/_rels/workbook.xml.rels":
                    oz.writestr(name, rels)
                else:
                    oz.writestr(name, data)
            oz.writestr("xl/vbaProject.bin", vba_bin)
        return True
    except Exception:
        return False


def enable_track_changes(file_path: Path) -> bool:
    """用本机 Excel 开启共享工作簿的"修订跟踪"。

    开启后，用户在本地 Excel 的修改会记录修改人、时间、单元格与旧值/新值。
    Excel 未安装、文件被占用或自动化失败时返回 False。
    """
    path = Path(file_path)
    temp_path = path.with_name(path.stem + "_tracktmp.xlsm")
    try:
        import win32com.client as win32  # type: ignore
    except Exception:
        return False

    for attempt in range(2):
        app = None
        workbook = None
        try:
            if temp_path.exists():
                temp_path.unlink()
            app = win32.DispatchEx("Excel.Application")
            app.Visible = False
            app.DisplayAlerts = False
            app.EnableEvents = False
            try:
                app.AutomationSecurity = 3
            except Exception:
                pass

            workbook = app.Workbooks.Open(str(path))
            workbook.TrackRevisions = True
            try:
                workbook.HighlightChangesOnScreen = True
            except Exception:
                pass
            workbook.SaveAs(str(temp_path), FileFormat=52, AccessMode=2)
            workbook.Close(SaveChanges=False)
            workbook = None
            app.Quit()
            app = None
            temp_path.replace(path)
            return True
        except Exception:
            if attempt == 0:
                time.sleep(1)
                continue
            return False
        finally:
            if workbook is not None:
                try:
                    workbook.Close(SaveChanges=False)
                except Exception:
                    pass
            if app is not None:
                try:
                    app.Quit()
                except Exception:
                    pass
            try:
                if temp_path.exists():
                    temp_path.unlink()
            except Exception:
                pass
    return False


class Exporter:
    """负责把源数据 + 人工字段渲染成 INV1/PKL1 工作簿。"""

    def __init__(self, template_dir: Path | None = None, price_file: Path | None = None) -> None:
        self.template_dir = Path(template_dir or TEMPLATE_DIR)
        self.prices = load_prices(price_file)
        self.transport_map = dict(TRANSPORT_TEMPLATES)
        self.consignee_data: dict[str, dict[str, Any]] = {}
        self._load_consignee_from_templates()

    # ---- 模板辅助 ----

    def _load_consignee_from_templates(self) -> None:
        mapping = {"空运模板.xlsx": "EASTEX", "海运模板.xlsx": "VANCO"}
        for tpl_name in mapping:
            tpl_path = self.template_dir / tpl_name
            if not tpl_path.exists():
                continue
            try:
                wb = openpyxl.load_workbook(str(tpl_path))
                ws = wb["INV1"]
                name = ws["B9"].value
                if not name:
                    wb.close()
                    continue
                lines = []
                for r in range(10, 16):
                    v = ws[f"B{r}"].value
                    if v is not None and str(v).strip():
                        lines.append(str(v).strip())
                wb.close()
                address, tel, email, vat = [], "", "", ""
                for line in lines:
                    up = line.upper()
                    if up.startswith("TEL"):
                        tel = line
                    elif up.startswith("EMAIL"):
                        email = line
                    elif up.startswith("VAT"):
                        vat = line
                    else:
                        address.append(line)
                self.consignee_data[str(name).strip()] = {
                    "name": str(name).strip(),
                    "address": address,
                    "tel": tel,
                    "email": email,
                    "vat": vat,
                }
            except Exception:
                continue

    def _write_header(self, ws, header: Optional[dict[str, str]]) -> None:
        if not header or not header.get("company_name"):
            return

        for m in list(ws.merged_cells.ranges):
            if m.min_row <= 4 and m.max_row >= 1 and m.min_col <= 1 and m.max_col >= 1:
                ws.unmerge_cells(str(m))

        ws["A1"] = header["company_name"]
        font1 = copy(ws["A1"].font)
        font1.bold = True
        font1.size = 18
        ws["A1"].font = font1
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

        address_phone = header.get("address_phone", "")
        if address_phone:
            ws["A2"] = address_phone
            font2 = copy(ws["A2"].font)
            font2.bold = True
            font2.size = 16
            ws["A2"].font = font2
            ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

        ws["A3"] = ""
        ws["A4"] = ""
        if ws["A3"].hyperlink:
            ws["A3"].hyperlink = None
        if ws["A4"].hyperlink:
            ws["A4"].hyperlink = None

        ws.merge_cells("A1:J1")
        ws.merge_cells("A2:J2")

    # ---- 主流程 ----

    def build(
        self,
        group_data: list[SourceData],
        manual: ManualFields | dict | None = None,
    ) -> BuildResult:
        """按 EXE 逻辑构建工作簿（不落盘）。

        Args:
            group_data: 同一 Delivery to 分组的源数据列表。
            manual: 人工填写字段；缺省视为空白（用于编辑页预填预览）。

        Returns:
            BuildResult，workbook 未保存。
        """
        fields = manual if isinstance(manual, ManualFields) else ManualFields.from_dict(manual)
        first = group_data[0]
        delivery_to = first.delivery_to or ""
        ship_mode = first.ship_mode or ""

        if "EASTEX" in delivery_to.upper():
            company_for_price, suffix = "Eastex", "-BT"
        elif "VANCO" in delivery_to.upper():
            company_for_price, suffix = "Vanco", "-CT"
        else:
            company_for_price, suffix = "Eastex", ""

        if "Air" in ship_mode:
            transport_suffix = "AIR"
        elif "Truck" in ship_mode:
            transport_suffix = "TRUCK"
        else:
            transport_suffix = "SEA"

        transport = get_transport_type(ship_mode, first.file_name) or "海运"

        supplier = (fields.supplier or "UNKNOWN").strip()
        bill_no = fields.bill_no.strip()
        file_stem = f"{supplier} INV#{bill_no}{suffix}({transport_suffix})"

        template_file = self.template_dir / self.transport_map.get(transport, "")
        if not template_file.exists():
            result = BuildResult(workbook=Workbook())
            result.ok = False
            result.company = company_for_price
            result.transport = transport
            result.bill_no = bill_no
            result.file_stem = file_stem
            result.messages.append(f"模板文件不存在: {template_file.name}")
            return result

        wb_out = openpyxl.load_workbook(str(template_file))
        result = BuildResult(workbook=wb_out)
        result.company = company_for_price
        result.transport = transport
        result.bill_no = bill_no
        result.file_stem = file_stem
        ws_inv1 = wb_out["INV1"]
        self._write_header(ws_inv1, first.header)

        # 动态定位 Total 行与 CONTAINER 行
        total_row = 37
        container_row = None
        for r in range(24, 60):
            c_val = ws_inv1[f"C{r}"].value
            a_val = ws_inv1[f"A{r}"].value
            if c_val and "Total" in str(c_val):
                total_row = r
            if a_val and "CONTAINER" in str(a_val).upper():
                container_row = r
        ref_default_row = total_row - 1
        result.total_row = total_row
        result.container_row = container_row

        # ---- 人工字段：Date / Invoice No / Ship mode ----
        # 只在有值时写入，与 EXE 一致：无 manual_data 时保留模板回退值
        if fields.date:
            ws_inv1["J9"] = fields.date
        if fields.bill_no:
            ws_inv1["J10"] = fields.bill_no
            ws_inv1["J10"].alignment = Alignment(horizontal="left", vertical="top")
        if fields.ship_mode:
            ws_inv1["B17"] = fields.ship_mode

        # 起运地（三种运输都写 A18，保留模板标签；值已含标签时直接写入避免重复）
        if fields.place_receipt:
            prefix, _ = _split_label(ws_inv1["A18"].value)
            value = fields.place_receipt.strip()
            if prefix.strip() and value.lower().startswith(prefix.strip().lower()):
                ws_inv1["A18"] = value
            else:
                ws_inv1["A18"] = f"{prefix}{value}"

        # 陆运：装货港写 A19；清掉模板残留的空格
        if transport == "陆运":
            if fields.port_loading:
                value = fields.port_loading.strip()
                if value.lower().startswith("port of loading:"):
                    ws_inv1["A19"] = value
                else:
                    ws_inv1["A19"] = f"Port of loading: {value}"
            if ws_inv1["F10"].value == " ":
                ws_inv1["F10"] = None

        # Consignee
        if delivery_to and delivery_to in self.consignee_data:
            consignee = self.consignee_data[delivery_to]
            for r in range(9, 16):
                ws_inv1[f"B{r}"] = None
            ws_inv1["B9"] = consignee["name"]
            for i, addr_line in enumerate(consignee["address"]):
                ws_inv1[f"B{10 + i}"] = addr_line
            addr_count = len(consignee["address"])
            ws_inv1[f"B{10 + addr_count}"] = consignee["tel"]
            ws_inv1[f"B{11 + addr_count}"] = consignee["email"]
            ws_inv1[f"B{12 + addr_count}"] = consignee["vat"]

        all_po: list[dict[str, Any]] = []
        all_bill_nos: list[str] = []
        for data in group_data:
            all_po.extend(data.po_data)
            if data.bill_no:
                all_bill_nos.append(str(data.bill_no).strip())

        if not all_po:
            result.ok = False
            result.messages.append("无面料数据")
            wb_out.close()
            return result

        net_weight_sum = sum(po["net_weight"] for po in all_po)
        ratio = weight_ratio(ship_mode)
        customs_weight = int(net_weight_sum * ratio)
        result.customs_weight = customs_weight

        fabric_type_data: dict[str, dict[str, Any]] = defaultdict(
            lambda: {"weight": 0.0, "descriptions": defaultdict(float)}
        )
        for po in all_po:
            ft = po["fabric_type"] or "UNKNOWN"
            fabric_type_data[ft]["weight"] += po["net_weight"]
            fabric_type_data[ft]["descriptions"][po["description"]] += po["net_weight"]
        sorted_ft = sorted(fabric_type_data.items(), key=lambda x: -x[1]["weight"])

        saved_merges = []
        ref_merge_restore: Optional[str] = None
        for m in list(ws_inv1.merged_cells.ranges):
            if m.min_row >= 24 and m.max_row <= total_row - 1:
                top_cell = ws_inv1.cell(row=m.min_row, column=m.min_col)
                if top_cell.value and "REF:" in str(top_cell.value):
                    # 清掉模板 REF 内容但保留合并结构，与模板表格线一致
                    ws_inv1.unmerge_cells(str(m))
                    top_cell.value = None
                    ref_merge_restore = str(m)
                    continue
                saved_merges.append(
                    (str(m), m.min_row, m.min_col, m.max_row, m.max_col)
                )
                ws_inv1.unmerge_cells(str(m))

        for r in range(23, total_row):
            for col in "ABCDEFGHIJK":
                cell = ws_inv1[f"{col}{r}"]
                if not isinstance(cell, MergedCell):
                    cell.value = None

        ws_inv1["A24"] = "N/M"

        desc_font = copy(ws_inv1["D25"].font)
        desc_alignment = copy(ws_inv1["D25"].alignment)
        desc_border = copy(ws_inv1["D25"].border)
        desc_fill = copy(ws_inv1["D25"].fill)
        desc_number_format = ws_inv1["D25"].number_format

        cur_row = 24
        total_amount = 0.0
        price_fill = PatternFill(fill_type="solid", start_color="FFF2CC")
        price_missing_fill = PatternFill(fill_type="solid", start_color="FFF8CBAD")
        for ft, fdata in sorted_ft:
            ft_weight = fdata["weight"]
            ft_descs = fdata["descriptions"]
            customs_w = int(ft_weight * ratio)
            # 单价：优先使用手动填写的覆盖值，否则按单价表自动读取
            override = fields.unit_prices.get(ft.strip().upper())
            if override is not None:
                ft_price = float(override)
            else:
                ft_price = lookup_price(self.prices, company_for_price, ft)
            ft_amount = round(ft_weight * ratio * ft_price, 2)
            total_amount += ft_amount

            ws_inv1[f"D{cur_row}"] = ft.upper()
            ws_inv1[f"G{cur_row}"] = customs_w
            ws_inv1[f"H{cur_row}"] = "KGS"
            ws_inv1[f"I{cur_row}"] = ft_price
            ws_inv1[f"J{cur_row}"] = ft_amount
            for col in ["D", "G", "H", "I", "J"]:
                src = ws_inv1[f"{col}24"]
                dst = ws_inv1[f"{col}{cur_row}"]
                if src.has_style:
                    dst.font = copy(src.font)
                    dst.alignment = copy(src.alignment)
                    dst.border = copy(src.border)
                    dst.fill = copy(src.fill)
                    dst.number_format = src.number_format
            if ft_price == 0.0 and override is None:
                ws_inv1[f"I{cur_row}"].fill = price_missing_fill
                result.messages.append(
                    f"面料 {ft.upper()} 未在单价表中找到单价，"
                    f"请在 INV1 表 I{cur_row} 单元格（橙色底纹）手动填写单价"
                )
            else:
                ws_inv1[f"I{cur_row}"].fill = price_fill
            cur_row += 1

            for desc in ft_descs:
                cell = ws_inv1[f"D{cur_row}"]
                cell.value = desc
                cell.font = copy(desc_font)
                cell.alignment = copy(desc_alignment)
                cell.border = copy(desc_border)
                cell.fill = copy(desc_fill)
                cell.number_format = desc_number_format
                cur_row += 1

        result.last_data_row = cur_row
        ref_row = max(cur_row, ref_default_row)
        result.ref_row = ref_row
        ref_value = " ".join(all_bill_nos)
        for coord, *_ in saved_merges:
            try:
                ws_inv1.merge_cells(coord)
            except Exception:
                pass
        if ref_merge_restore:
            # 恢复模板 REF 区域的合并结构（内容已清空），表格线按模板来
            try:
                ws_inv1.merge_cells(ref_merge_restore)
            except Exception:
                pass
        ws_inv1[f"A{ref_row}"] = f"REF:{ref_value}"

        # 集装箱区域：清空 A/B 列，写入人工输入（最多 3 行）
        if container_row:
            for r in range(container_row + 1, container_row + 10):
                for col in ["A", "B"]:
                    cell = ws_inv1[f"{col}{r}"]
                    if not isinstance(cell, MergedCell):
                        cell.value = None
            if fields.container_info:
                for i, line in enumerate(fields.container_info.split("\n")[:3]):
                    ws_inv1[f"A{container_row + 1 + i}"] = line

        # 复制源 InputTable（保留格式），与 EXE 输出结构一致
        if "InputTable" in wb_out.sheetnames:
            del wb_out["InputTable"]
        for i, data in enumerate(group_data):
            inv_no = str(data.bill_no).strip() if data.bill_no else f"file{i}"
            try:
                with warnings_shield():
                    src_wb = openpyxl.load_workbook(str(data.file_path), data_only=True)
            except Exception as exc:
                result.messages.append(f"读取源文件失败 {data.file_name}: {exc}")
                continue
            try:
                if "InputTable" in src_wb.sheetnames:
                    target = "InputTable" if i == 0 else f"InputTable_{inv_no}"
                    copy_sheet_with_format(src_wb, "InputTable", wb_out, target)
            finally:
                src_wb.close()

        # ---- PKL1 ----
        total_rolls = sum(d.rolls for d in group_data)
        total_gross = sum(d.gross_weight for d in group_data)
        if "PKL1" in wb_out.sheetnames:
            ws_pkl = wb_out["PKL1"]
            pkl_ref_row = None
            for r in range(24, 55):
                v = ws_pkl[f"A{r}"].value
                if v and "REF" in str(v):
                    pkl_ref_row = r
                    break

            ws_pkl["A24"] = total_rolls
            ws_pkl["B24"] = "PACKS"
            ws_pkl["K24"] = total_gross

            for r in range(24, cur_row):
                ws_pkl[f"E{r}"] = f"='INV1'!D{r}"

            if transport == "空运":
                ws_pkl["I17"] = fields.flight_no or first.vessel_flight or None

            if pkl_ref_row:
                for r in range(25, pkl_ref_row):
                    for col in ["A", "B", "H", "I", "J", "K"]:
                        cell = ws_pkl[f"{col}{r}"]
                        if not isinstance(cell, MergedCell):
                            cell.value = None
                for r in range(cur_row, pkl_ref_row):
                    cell = ws_pkl[f"E{r}"]
                    if not isinstance(cell, MergedCell):
                        cell.value = None
                ws_pkl[f"A{pkl_ref_row}"] = f"REF:{ref_value}"

                clear_from = pkl_ref_row + 1
                clear_to = pkl_ref_row + 8
                for row in range(clear_from, clear_to):
                    cell_a = ws_pkl[f"A{row}"]
                    cell_h = ws_pkl[f"H{row}"]
                    if cell_a.value and str(cell_a.value).startswith("='INV1'"):
                        cell_a.value = None
                    if cell_h.value and str(cell_h.value).startswith("='INV1'"):
                        cell_h.value = None

        result.total_amount = round(total_amount, 2)
        result.messages.append(
            f"提单号 {ref_value} | 报关重量 {customs_weight}KG | 金额 ${total_amount:,.2f}"
        )
        return result

    # ---- 落盘 ----

    def save_xlsm(
        self,
        workbook: Workbook,
        out_path: Path,
        carrier: Path,
        keep_xlsx: bool = False,
    ) -> Path:
        """保存工作簿并注入 VBA，返回最终 .xlsm 路径。"""
        out_path = Path(out_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        temp_file = out_path.parent / (out_path.stem + "_vbatmp.xlsx")
        workbook.save(str(temp_file))
        workbook.close()

        if out_path.exists():
            try:
                out_path.unlink()
            except PermissionError:
                pass

        if inject_vba(temp_file, Path(carrier), out_path):
            if not keep_xlsx:
                temp_file.unlink(missing_ok=True)
        else:
            temp_file.replace(out_path)
        return out_path


class warnings_shield:
    """屏蔽 openpyxl 读取旧文件时的 Data Validation 告警。"""

    def __enter__(self):
        import warnings

        self._ctx = warnings.catch_warnings()
        self._ctx.__enter__()
        warnings.simplefilter("ignore")
        return self

    def __exit__(self, *exc):
        return self._ctx.__exit__(*exc)
