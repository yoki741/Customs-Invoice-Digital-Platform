# -*- coding: utf-8 -*-
"""装柜列表（Shipment List）生成。

把一次导出的发票汇总成装柜列表模板中的一行：
- 表头固定在模板第 1 行，数据从第 2 行开始
- ETA 按运输方式自动推算：空运 ETD+1，海运/汽运 ETD+8
- VESSEL / Carrier 有源数据时预填，否则留空供人工填写
"""

from __future__ import annotations

import re
from copy import copy
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Optional, Sequence

import openpyxl

from .config import SHIPMENT_LIST_DIR, SHIPMENT_LIST_TEMPLATE
from .source_reader import SourceData

SHEET_NAME = "Sheet1"
FIRST_DATA_ROW = 2
EXPECTED_COLS = 24

_PLACE_PREFIXES = (
    "port of loading:",
    "port of loading：",
    "place of receipt:",
    "place of receipt：",
    "origin:",
    "origin：",
    "装货港:",
    "装货港：",
    "起运地:",
    "起运地：",
)


@dataclass
class ShipmentRow:
    """模板中的一行装柜记录。"""

    date: Optional[date]
    fty_inv: str = ""
    shipper: str = ""
    est_vco: str = ""
    custom_no: str = ""
    package: Optional[float] = None
    blading: str = ""
    containers: list[str] = field(default_factory=list)
    nw_kg: float = 0.0
    gw_kg: float = 0.0
    description: str = ""
    mode: str = ""
    etd: Optional[date] = None
    eta: Optional[date] = None
    pol: str = ""
    vessel: str = ""
    carrier: str = ""
    container_type: str = ""
    cbm: Optional[float] = None
    machine_cabinet: str = ""
    aux_cabinet: str = ""

    def cells(self) -> list[Any]:
        """按模板 A-X 列返回单元格值。"""
        values: list[Any] = [
            self.date,
            self.fty_inv,
            self.shipper,
            self.est_vco,
            self.custom_no,
            self.package,
            self.blading,
        ]
        boxes = list(self.containers[:4])
        values.extend(boxes + [None] * (4 - len(boxes)))
        values.extend(
            [
                self.nw_kg,
                self.gw_kg,
                self.description,
                self.mode,
                self.etd,
                self.eta,
                self.pol,
                self.vessel,
                self.carrier,
                self.container_type,
                self.cbm,
                self.machine_cabinet,
                self.aux_cabinet,
            ]
        )
        return values

    def as_dict(self) -> dict[str, Any]:
        """供数据库 / 页面使用的扁平字典。"""
        return {
            "date": _iso(self.date),
            "fty_inv": self.fty_inv,
            "shipper": self.shipper,
            "est_vco": self.est_vco,
            "custom_no": self.custom_no,
            "package": self.package,
            "blading": self.blading,
            "container_1": self.containers[0] if len(self.containers) > 0 else "",
            "container_2": self.containers[1] if len(self.containers) > 1 else "",
            "container_3": self.containers[2] if len(self.containers) > 2 else "",
            "container_4": self.containers[3] if len(self.containers) > 3 else "",
            "nw_kg": self.nw_kg,
            "gw_kg": self.gw_kg,
            "description": self.description,
            "mode": self.mode,
            "etd": _iso(self.etd),
            "eta": _iso(self.eta),
            "pol": self.pol,
            "vessel": self.vessel,
            "carrier": self.carrier,
            "container_type": self.container_type,
            "cbm": self.cbm,
            "machine_cabinet": self.machine_cabinet,
            "aux_cabinet": self.aux_cabinet,
        }


def _iso(value: Optional[date]) -> str:
    return value.strftime("%Y-%m-%d") if value else ""


def parse_date(value: Any) -> Optional[date]:
    """把日期 / 日期时间 / 常见字符串统一转成 date。"""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    candidates = {text, text[:10], text[:19]}
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d", "%m/%d/%Y"):
        for candidate in candidates:
            try:
                return datetime.strptime(candidate, fmt).date()
            except ValueError:
                continue
    return None


def _text(value: Any) -> str:
    """清理单元格文本：去掉 Excel 换行标记与多余空白。"""
    text = str(value or "")
    text = text.replace("_x000D_", " ").replace("\r", " ").replace("\n", " ")
    return re.sub(r"\s+", " ", text).strip()


def _num(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def company_from(sources: Sequence[SourceData]) -> str:
    """从 Delivery to / Bill to 识别我司工厂（Eastex / Vanco）。"""
    for source in sources:
        blob = f"{source.delivery_to} {source.bill_to}".upper()
        if "VANCO" in blob:
            return "Vanco"
        if "EASTEX" in blob:
            return "Eastex"
    return ""


def suffix_for(company: str) -> str:
    """清关文件发票号后缀：Eastex=_BT，Vanco=_CT。"""
    return "CT" if company == "Vanco" else "BT"


def mode_from(transport: str = "", ship_mode: str = "") -> str:
    """运输方式统一成模板使用的 AIR / SEA / TRUCK。"""
    blob = f"{transport} {ship_mode}".upper()
    if "AIR" in blob or "空运" in blob:
        return "AIR"
    if "TRUCK" in blob or "陆运" in blob:
        return "TRUCK"
    return "SEA"


def eta_for(mode: str, etd: Optional[date]) -> Optional[date]:
    """ETA 规则：空运 ETD+1，海运/汽运 ETD+8。"""
    if etd is None:
        return None
    return etd + timedelta(days=1 if mode == "AIR" else 8)


def _distinct(values: Sequence[Any]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        text = _text(value)
        if text and text not in seen:
            seen.add(text)
            result.append(text)
    return result


def _clean_place(value: str) -> str:
    text = value.strip()
    lowered = text.lower()
    for prefix in _PLACE_PREFIXES:
        if lowered.startswith(prefix):
            return text[len(prefix) :].strip()
    return text


def _containers(manual_info: str, source_numbers: Sequence[str]) -> list[str]:
    """柜号逐格对应 Container 1-4：格子里填什么就展示什么，不做抽取改写。

    手工填写的 container_info 按行拆分，第 N 行对应报关清单
    CONTAINER NO / SEAL NO: 往下第 N 个单元格（Container N）；
    没有手工填写时依次取源文件柜号。
    """
    if manual_info:
        cells = [str(line or "").strip() for line in manual_info.split("\n")[:4]]
    else:
        cells = [_text(number) for number in source_numbers[:4]]
    while cells and not cells[-1]:
        cells.pop()
    return cells


def build_row(
    sources: Sequence[SourceData],
    manual: Optional[dict[str, Any]] = None,
    *,
    company: str = "",
    transport: str = "",
) -> ShipmentRow:
    """把一批（同一次导出的）源文件汇总成一行装柜记录。"""
    if not sources:
        raise ValueError("没有可汇总的源文件")
    manual = manual or {}
    first = sources[0]

    resolved_company = company or company_from(sources)
    bill_no = _text(manual.get("bill_no")) or _text(first.invoice_no)
    mode = mode_from(transport, first.ship_mode or _text(manual.get("ship_mode")))

    etd = parse_date(first.invoice_date) or parse_date(first.etd)
    row_date = parse_date(manual.get("date")) or etd

    po_numbers = _distinct(
        [item.po_no for source in sources for item in source.items]
    )
    invoice_numbers = _distinct([source.invoice_no for source in sources])
    shippers = _distinct([source.supplier_code for source in sources])
    agents = _distinct([source.customs_agent for source in sources])
    vessels = _distinct([source.vessel_flight for source in sources])
    descriptions = _distinct(
        [item.description for source in sources for item in source.items]
    )
    if not descriptions:
        descriptions = _distinct(
            [item.fabric_type for source in sources for item in source.items]
        )

    rolls = sum(_num(source.rolls) for source in sources)
    if rolls <= 0:
        rolls = sum(
            _num(item.carton) for source in sources for item in source.items
        )

    package = round(rolls, 3)
    if package == int(package):
        package = int(package)

    place = _clean_place(_text(manual.get("place_receipt"))) or _text(
        first.origin_place
    )

    return ShipmentRow(
        date=row_date,
        fty_inv=", ".join(po_numbers) or ", ".join(invoice_numbers),
        shipper=", ".join(shippers),
        est_vco=resolved_company,
        custom_no=f"{bill_no}_{suffix_for(resolved_company)}" if bill_no else "",
        package=package or None,
        blading=bill_no,
        containers=_containers(
            str(manual.get("container_info") or ""),
            [source.container_number or "" for source in sources],
        ),
        nw_kg=round(sum(_num(source.net_weight) for source in sources), 3),
        gw_kg=round(sum(_num(source.gross_weight) for source in sources), 3),
        description=" / ".join(descriptions),
        mode=mode,
        etd=etd,
        eta=eta_for(mode, etd),
        pol=place,
        vessel=_text(manual.get("flight_no")) or ", ".join(vessels),
        carrier=", ".join(agents),
    )


def _company_from_consignee(value: str) -> str:
    """从清关文件 Consignee 名称识别我司工厂（Vanco / Eastex）。"""
    blob = value.upper()
    if "VANCO" in blob:
        return "Vanco"
    if "EASTEX" in blob or "DAKOTA" in blob:
        return "Eastex"
    return ""


def _read_ref(ws: Any) -> str:
    """读取 INV1 的 REF 行内容，去掉 "REF:" 前缀。"""
    for row in range(24, 61):
        value = ws[f"A{row}"].value
        if value is None:
            continue
        text = str(value).strip()
        if text.upper().startswith("REF"):
            idx = text.find(":")
            return text[idx + 1 :].strip() if idx >= 0 else text[3:].strip()
    return ""


def _read_fabric_types(ws: Any) -> list[str]:
    """读取 INV1 的面料分类（D 列中带报关重量 G 的行）。"""
    found: list[str] = []
    for row in range(24, 61):
        d_val = ws[f"D{row}"].value
        g_val = ws[f"G{row}"].value
        if d_val is None or g_val is None:
            continue
        text = _text(d_val)
        if text and text not in found:
            found.append(text)
    return found


def _read_containers(ws: Any) -> list[str]:
    """读取 INV1 的 CONTAINER NO / SEAL NO: 往下第 1-4 格。

    格子里填什么，装柜列表 Container 1-4 就逐个展示什么样（不抽取柜号）。
    """
    container_row: Optional[int] = None
    for row in range(1, 61):
        value = ws[f"A{row}"].value
        if value and "CONTAINER" in str(value).upper():
            container_row = row
            break
    if container_row is None:
        return []
    cells = [_text(ws[f"A{container_row + 1 + i}"].value) for i in range(4)]
    while cells and not cells[-1]:
        cells.pop()
    return cells


def _read_packages(ws: Any) -> Optional[float]:
    """读取 PKL 中 PACKS 前面的数字（如 307 PACKS → 307）。

    优先取同一单元格文本里紧挨 "packs" 前的数字，否则取同一行 PACKS
    左侧最近的数值单元格；出现多行 PACKS 时求和。找不到返回 None。
    """
    total = 0.0
    found = False
    for row in ws.iter_rows():
        for cell in row:
            text = _text(cell.value)
            if "PACKS" not in text.upper():
                continue
            match = re.search(r"([\d][\d.,]*)\s*PACKS", text, re.IGNORECASE)
            if match:
                total += _num(match.group(1).replace(",", ""))
                found = True
                continue
            for col in range(cell.column - 1, 0, -1):
                value = ws.cell(row=cell.row, column=col).value
                if value in (None, ""):
                    continue
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    total += float(value)
                    found = True
                break
    return total if found else None


def build_row_from_result(
    workbook: openpyxl.Workbook,
    *,
    company: str = "",
    transport: str = "",
    customs_weight: float = 0.0,
) -> ShipmentRow:
    """从清关文件工作簿（INV1/PKL1）提取装柜列表一行。

    数据来源是导出生成的清关文件（结果文件），而非源发票文件：
    - 抬头（Shipper）= INV1!A1，Consignee（EST/VCO）= INV1!B9
    - Date = INV1!J9，Invoice No = INV1!J10，Ship mode = INV1!B17
    - Fty Inv = INV1 REF 行，POL = INV1!A18（陆运 A19）
    - 分类 = INV1 D 列，集装箱 = INV1 CONTAINER 区
    - 净重 = 报关净重（传入 customs_weight），毛重 = PKL1!K24
    - Package = PKL1 中 PACKS 前面的数字
    - ETD / VESSEL / Carrier / 柜型 / CBM / 机器柜 / 辅料柜 为手动字段，留空供人工补齐
      （机器柜 / 辅料柜仅在 Shipment List 页面行内编辑，不随导出写入）
    """
    ws_inv1 = workbook["INV1"]

    shipper = _text(ws_inv1["A1"].value)
    consignee = _text(ws_inv1["B9"].value)
    resolved_company = company or _company_from_consignee(consignee)
    row_date = parse_date(ws_inv1["J9"].value)
    bill_no = _text(ws_inv1["J10"].value)
    mode = mode_from(transport, _text(ws_inv1["B17"].value))

    fty_inv = _read_ref(ws_inv1)
    pol_cell = "A19" if transport == "陆运" else "A18"
    pol = _clean_place(_text(ws_inv1[pol_cell].value))
    descriptions = _read_fabric_types(ws_inv1)
    containers = _read_containers(ws_inv1)

    gw_kg = 0.0
    package: Optional[float] = None
    if "PKL1" in workbook.sheetnames:
        ws_pkl = workbook["PKL1"]
        gw_kg = _num(ws_pkl["K24"].value)
        package = _read_packages(ws_pkl)
    if package is not None:
        package = round(package, 3)
        if package == int(package):
            package = int(package)

    return ShipmentRow(
        date=row_date,
        fty_inv=fty_inv,
        shipper=shipper,
        est_vco=resolved_company,
        custom_no=f"{bill_no}_{suffix_for(resolved_company)}" if bill_no else "",
        package=package or None,
        blading=bill_no,
        containers=containers,
        nw_kg=round(_num(customs_weight), 3),
        gw_kg=round(gw_kg, 3),
        description=" / ".join(descriptions),
        mode=mode,
        etd=None,
        eta=None,
        pol=pol,
        vessel="",
        carrier="",
        container_type="",
        cbm=None,
        machine_cabinet="",
        aux_cabinet="",
    )


def month_of(value: Optional[date]) -> str:
    """返回 YYYY-MM；无日期时用当前月份。"""
    return (value or date.today()).strftime("%Y-%m")


def month_path(month: str) -> Path:
    """月度装柜列表文件路径（month 需为 YYYY-MM）。"""
    return SHIPMENT_LIST_DIR / f"Shipment List {month}.xlsx"


def _migrate_template_headers(ws: Any) -> None:
    """兼容旧版模板表头：ID 原在 W 列，让位给机器柜 / 辅料柜后移到 Y 列。

    模板文件本身已更新时此处无操作；模板被占用改不了时也能保证
    生成文件的表头与 cells() 列序（W=机器柜、X=辅料柜）一致。
    """
    if str(ws["W1"].value or "").strip() != "ID" or ws["Y1"].value:
        return
    src = ws["W1"]
    ws["W1"].value = "机器柜"
    ws["X1"].value = "辅料柜"
    ws["Y1"].value = "ID"
    for dst in (ws["X1"], ws["Y1"]):
        dst.font = copy(src.font)
        dst.border = copy(src.border)
        dst.fill = copy(src.fill)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format
    width = ws.column_dimensions["W"].width
    if width:
        ws.column_dimensions["X"].width = width
        ws.column_dimensions["Y"].width = width


def write_workbook(
    rows: Sequence[ShipmentRow],
    out_path: str | Path,
    template: str | Path | None = None,
) -> Path:
    """按模板生成装柜列表文件，返回输出路径。"""
    out = Path(out_path)
    source_template = Path(template or SHIPMENT_LIST_TEMPLATE)
    if not source_template.exists():
        raise FileNotFoundError(f"装柜列表模板不存在: {source_template}")

    workbook = openpyxl.load_workbook(str(source_template))
    try:
        ws = workbook[SHEET_NAME]
        _migrate_template_headers(ws)
        for row in range(FIRST_DATA_ROW, ws.max_row + 1):
            for col in range(1, EXPECTED_COLS + 1):
                ws.cell(row=row, column=col).value = None

        for offset, record in enumerate(rows):
            target_row = FIRST_DATA_ROW + offset
            for col, value in enumerate(record.cells(), start=1):
                cell = ws.cell(row=target_row, column=col)
                style_cell = ws.cell(row=FIRST_DATA_ROW, column=col)
                if target_row != FIRST_DATA_ROW:
                    cell.font = copy(style_cell.font)
                    cell.border = copy(style_cell.border)
                    cell.fill = copy(style_cell.fill)
                    cell.alignment = copy(style_cell.alignment)
                    cell.number_format = style_cell.number_format
                cell.value = value

        out.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(str(out))
    finally:
        workbook.close()
    return out
