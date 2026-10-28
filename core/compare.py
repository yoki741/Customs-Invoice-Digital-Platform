# -*- coding: utf-8 -*-
"""两份报关发票文件的逐格内容比对 + 文档属性提取。

指纹（SHA256）不一致时，用它定位「具体改了哪些单元格」。跳过内部追踪表
（_PLATFORM_TRACE / _LOG），只比对业务可见内容。

文档属性（docProps/core.xml + app.xml）由 Excel 在保存时自动更新，
可作为「谁、何时保存过该文件」的旁证。
"""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook

# 内部追踪表，不参与内容比对
SKIP_SHEETS = {"_PLATFORM_TRACE", "_LOG"}


def _xml_tag(text: str, tag: str) -> str:
    """取 XML 里第一个指定标签的文本（无则空串）。"""
    m = re.search(rf"<{tag}[^>]*>(.*?)</{tag}>", text, re.S)
    return m.group(1).strip() if m else ""


def read_doc_props(data: bytes) -> dict[str, str]:
    """提取文档属性（谁、何时保存过该文件）。

    Excel 保存时自动更新 lastModifiedBy / modified，可作旁证。
    返回 ``{creator, last_modified_by, created, modified, application}``，
    读取失败的项为空串。
    """
    props = {
        "creator": "",
        "last_modified_by": "",
        "created": "",
        "modified": "",
        "application": "",
    }
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            try:
                core = z.read("docProps/core.xml").decode("utf-8", "replace")
            except KeyError:
                core = ""
            try:
                app = z.read("docProps/app.xml").decode("utf-8", "replace")
            except KeyError:
                app = ""
    except Exception:  # noqa: BLE001
        return props
    if core:
        props["creator"] = _xml_tag(core, "dc:creator")
        props["last_modified_by"] = _xml_tag(core, "cp:lastModifiedBy")
        props["created"] = _xml_tag(core, "dcterms:created")
        props["modified"] = _xml_tag(core, "dcterms:modified")
    if app:
        props["application"] = _xml_tag(app, "Application")
    return props


def _norm(value: object) -> str | None:
    """把单元格值归一化为可比文本；空值返回 None。

    整型浮点归一化（123.0 → 123），避免 Excel 重存带来的 123 vs 123.0 误报。
    """
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _cell_map(workbook) -> dict[tuple[str, str], str]:
    """收集工作簿全部可见单元格 {(表名, 坐标): 文本值}。"""
    result: dict[tuple[str, str], str] = {}
    for ws in workbook.worksheets:
        if ws.title in SKIP_SHEETS:
            continue
        for row in ws.iter_rows():
            for cell in row:
                text = _norm(cell.value)
                if text is None:
                    continue
                result[(ws.title, cell.coordinate)] = text
    return result


def diff_xlsm(original_path: Path, modified_bytes: bytes) -> dict:
    """比对原始导出文件与回传文件，返回 ``{count, changes}``。

    ``changes`` 每项为 ``{sheet, cell, original, modified}``，original / modified
    为归一化文本（无改动方向则对应值为 None）。
    """
    original_wb = load_workbook(str(original_path), data_only=False, keep_vba=False)
    modified_wb = load_workbook(io.BytesIO(modified_bytes), data_only=False, keep_vba=False)
    try:
        original = _cell_map(original_wb)
        modified = _cell_map(modified_wb)
    finally:
        original_wb.close()
        modified_wb.close()

    changes: list[dict] = []
    for key in sorted(set(original) | set(modified)):
        old = original.get(key)
        new = modified.get(key)
        if old != new:
            changes.append(
                {"sheet": key[0], "cell": key[1], "original": old, "modified": new}
            )
    return {"count": len(changes), "changes": changes}
