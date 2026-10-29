# -*- coding: utf-8 -*-
"""读取源 .xlsm 的 InputTable / Inv 工作表。

字段坐标与 invoice_gui.InvoiceProcessor.read_source_file 完全一致，
并额外补全了明细各列（Colour / Size / Carton / 数量 / 单价等）供前端展示。
"""

from __future__ import annotations

import warnings
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

import openpyxl

# EXE 只读取 17-44 行（28 条明细），导出必须保持一致，见 Exporter
EXPORT_FIRST_ROW = 17
EXPORT_LAST_ROW = 44
# 源表 PO 行一直排到 66 行，明细展示读到 66
DETAIL_FIRST_ROW = 17
DETAIL_LAST_ROW = 66


@dataclass
class DetailItem:
    """InputTable 明细行。"""

    row: int
    line_no: Optional[str] = None
    po_no: Optional[str] = None
    description: Optional[str] = None
    colour: Optional[str] = None
    size: Optional[str] = None
    fabric_type: Optional[str] = None
    carton: Optional[float] = None
    net_weight: Optional[float] = None
    gross_weight: Optional[float] = None
    length: Optional[float] = None
    quantity: Optional[float] = None
    unit_price: Optional[float] = None
    amount: Optional[float] = None

    @property
    def in_export_range(self) -> bool:
        """是否落在 EXE 导出所使用的行区间。"""
        return EXPORT_FIRST_ROW <= self.row <= EXPORT_LAST_ROW


@dataclass
class SourceData:
    """单个源文件解析结果。"""

    file_path: Path
    file_name: str = ""

    # 抬头
    invoice_no: Optional[str] = None
    invoice_date: Optional[str] = None
    bill_to: Optional[str] = None
    delivery_to: str = ""
    ship_mode: Optional[str] = None
    ship_term: Optional[str] = None
    vessel_flight: str = ""
    etd: Optional[str] = None
    eta: Optional[str] = None
    container_number: Optional[str] = None
    seal_number: Optional[str] = None
    supplier_code: Optional[str] = None
    origin_place: Optional[str] = None
    origin_country: Optional[str] = None
    destination: Optional[str] = None
    currency: Optional[str] = None
    customs_agent: Optional[str] = None
    length_unit: Optional[str] = None

    # 汇总
    rolls: float = 0.0
    net_weight: float = 0.0
    gross_weight: float = 0.0
    total_length: float = 0.0
    total_lbs: float = 0.0
    total_amount: float = 0.0

    # Inv sheet 抬头
    header: dict[str, str] = field(default_factory=dict)

    items: list[DetailItem] = field(default_factory=list)

    def __post_init__(self) -> None:
        if not self.file_name:
            self.file_name = Path(self.file_path).name

    # ---- 供导出逻辑使用的兼容视图（与 invoice_gui.py 的数据结构一致） ----

    @property
    def bill_no(self) -> Optional[str]:
        """EXE 中的 bill_no 即源表 Invoice No.。"""
        return self.invoice_no

    @property
    def po_data(self) -> list[dict[str, Any]]:
        """EXE 使用的明细子集（description / fabric_type / net_weight）。"""
        rows = [i for i in self.items if i.in_export_range and i.description and i.net_weight]
        return [
            {
                "description": i.description,
                "fabric_type": (i.fabric_type or "").strip(),
                "net_weight": float(i.net_weight or 0),
            }
            for i in rows
        ]

    @property
    def export_items(self) -> list[DetailItem]:
        return [i for i in self.items if i.in_export_range and i.description and i.net_weight]


def _text(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _num(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _date_text(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip() or None


def read_source_file(file_path: str | Path) -> SourceData:
    """解析单个源 .xlsm 文件。

    Args:
        file_path: 源文件路径，必须含 InputTable 工作表。

    Returns:
        SourceData：抬头 + 汇总 + 明细。

    Raises:
        FileNotFoundError: 文件不存在。
        KeyError: 缺少 InputTable 工作表。
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(path)

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        # read_only 模式解析更快（只需取值，不需要样式），编辑页打开提速
        wb = openpyxl.load_workbook(
            str(path), data_only=True, read_only=True, keep_vba=False
        )

    try:
        if "InputTable" not in wb.sheetnames:
            raise KeyError(f"{path.name} is missing the InputTable worksheet")

        ws = wb["InputTable"]
        data = SourceData(file_path=path)

        data.supplier_code = _text(ws["C3"].value)
        data.invoice_no = _text(ws["C5"].value)
        data.invoice_date = _date_text(ws["C6"].value)
        data.vessel_flight = _text(ws["C7"].value) or ""
        data.etd = _date_text(ws["C8"].value)
        data.eta = _date_text(ws["C9"].value)
        data.container_number = _text(ws["C10"].value)
        data.seal_number = _text(ws["C11"].value)

        data.bill_to = _text(ws["G4"].value)
        data.delivery_to = _text(ws["G5"].value) or ""
        data.ship_mode = _text(ws["G6"].value)
        data.ship_term = _text(ws["G7"].value)
        data.destination = _text(ws["G9"].value)
        data.currency = _text(ws["G10"].value)
        data.customs_agent = _text(ws["G11"].value)
        data.length_unit = _text(ws["G12"].value)

        data.origin_place = _text(ws["E3"].value)
        data.origin_country = _text(ws["E4"].value)

        data.rolls = float(_num(ws["G14"].value) or 0)
        data.net_weight = float(_num(ws["H14"].value) or 0)
        data.gross_weight = float(_num(ws["I14"].value) or 0)
        data.total_length = float(_num(ws["J14"].value) or 0)
        data.total_lbs = float(_num(ws["K14"].value) or 0)
        data.total_amount = float(_num(ws["M14"].value) or 0)

        for row in range(DETAIL_FIRST_ROW, DETAIL_LAST_ROW + 1):
            desc = ws[f"C{row}"].value
            net = ws[f"H{row}"].value
            if not desc or not net:
                continue
            data.items.append(
                DetailItem(
                    row=row,
                    line_no=_text(ws[f"A{row}"].value),
                    po_no=_text(ws[f"B{row}"].value),
                    description=_text(desc),
                    colour=_text(ws[f"D{row}"].value),
                    size=_text(ws[f"E{row}"].value),
                    fabric_type=_text(ws[f"F{row}"].value),
                    carton=_num(ws[f"G{row}"].value),
                    net_weight=_num(net),
                    gross_weight=_num(ws[f"I{row}"].value),
                    length=_num(ws[f"J{row}"].value),
                    quantity=_num(ws[f"K{row}"].value),
                    unit_price=_num(ws[f"L{row}"].value),
                    amount=_num(ws[f"M{row}"].value),
                )
            )

        if "Inv" in wb.sheetnames:
            inv_ws = wb["Inv"]
            data.header = {
                "company_name": str(inv_ws["B1"].value or "").strip(),
                "address_phone": str(inv_ws["B2"].value or "").strip(),
            }

        return data
    finally:
        wb.close()
