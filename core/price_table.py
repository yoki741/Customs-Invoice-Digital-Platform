# -*- coding: utf-8 -*-
"""单价表读取与公司 / 运输方式判定。"""

from __future__ import annotations

from pathlib import Path
from typing import Optional

import openpyxl

from .config import PRICE_FILE

# 单价表固定行区间（与 invoice_gui.py 保持一致）
_EASTEX_ROWS = range(2, 19)
_VANCO_ROWS = range(21, 38)

WEIGHT_RATIOS = {"By Sea": 0.94, "By Air": 0.95, "By Truck": 0.94}


def load_prices(price_file: Path | None = None) -> dict[str, dict[str, float]]:
    """读取 Eastex / Vanco 的 Fabric Type 单价。"""
    path = price_file or PRICE_FILE
    prices: dict[str, dict[str, float]] = {"Eastex": {}, "Vanco": {}}
    try:
        wb = openpyxl.load_workbook(str(path))
        ws = wb["Sheet1"]
        for row in _EASTEX_ROWS:
            ft, up = ws[f"B{row}"].value, ws[f"C{row}"].value
            if ft and up:
                prices["Eastex"][str(ft).strip().lower()] = float(up)
        for row in _VANCO_ROWS:
            ft, up = ws[f"B{row}"].value, ws[f"C{row}"].value
            if ft and up:
                prices["Vanco"][str(ft).strip().lower()] = float(up)
        wb.close()
    except Exception:
        pass
    return prices


def lookup_price(
    prices: dict[str, dict[str, float]], company: Optional[str], fabric_type: Optional[str]
) -> float:
    """按 Fabric Type 查单价，支持模糊包含匹配。"""
    if not fabric_type or not prices:
        return 0.0
    table = prices.get(company or "", {})
    key = fabric_type.strip().lower()
    if key in table:
        return table[key]
    for price_key, value in table.items():
        if key in price_key or price_key in key:
            return value
    return 0.0


def get_company_name(bill_to: Optional[str]) -> Optional[str]:
    """由 Bill to / Delivery to 判定归属公司。"""
    if not bill_to:
        return None
    upper = bill_to.upper()
    if "DAKOTA" in upper:
        return "Eastex"
    if "VANCO" in upper:
        return "Vanco"
    return None


def get_transport_type(ship_mode: Optional[str], file_name: str = "") -> Optional[str]:
    """判定运输方式（海运 / 空运 / 陆运）。"""
    if ship_mode:
        if "Sea" in ship_mode:
            return "海运"
        if "Air" in ship_mode:
            return "空运"
        if "Truck" in ship_mode:
            return "陆运"
    upper = (file_name or "").upper()
    if "SEA" in upper:
        return "海运"
    if "AIR" in upper:
        return "空运"
    if "TRUCK" in upper:
        return "陆运"
    return None


def weight_ratio(ship_mode: Optional[str]) -> float:
    """报关重量换算比例。"""
    ratio = 0.94
    if not ship_mode:
        return ratio
    for key, value in WEIGHT_RATIOS.items():
        if key.lower() in ship_mode.lower():
            return value
    return ratio
