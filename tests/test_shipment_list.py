# -*- coding: utf-8 -*-
"""装柜列表（Shipment List）映射、ETA 与模板写入测试。"""

from __future__ import annotations

import sys
import tempfile
import warnings
from datetime import date
from pathlib import Path

warnings.simplefilter("ignore")
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402

from core.shipment_list import (  # noqa: E402
    ShipmentRow,
    _containers,
    _read_packages,
    build_row,
    build_row_from_result,
    eta_for,
    month_of,
    mode_from,
    parse_date,
    write_workbook,
)
from core.source_reader import DetailItem, SourceData  # noqa: E402


def _source(**kwargs) -> SourceData:
    defaults = dict(
        file_path=Path("sample.xlsm"),
        invoice_no="SHSJ2610051A",
        invoice_date="2026-10-21",
        supplier_code="LEFEI",
        delivery_to="VANCO INDUSTRIAL CO., LTD",
        ship_mode="By Sea",
        vessel_flight="SITC RUNDE V.2622S",
        container_number="FBLU0186990",
        origin_place="NINGBO, CHINA",
        customs_agent="SITC",
        rolls=307,
        net_weight=6664,
        gross_weight=6900,
    )
    po_no = kwargs.pop("po_no", "26LF1025-1")
    defaults.update(kwargs)
    data = SourceData(**defaults)
    data.items = [
        DetailItem(
            row=17,
            po_no=po_no,
            description="KNITTED FABRIC",
            fabric_type="JC",
            carton=10,
            net_weight=100,
        )
    ]
    return data


def _result_workbook() -> openpyxl.Workbook:
    """构造一个最小清关文件工作簿（INV1/PKL1）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "INV1"
    ws["A1"] = "ZHEJIANG LEFEI TEXTILE CO.,LTD."
    ws["B9"] = "VANCO INDUSTRIAL CO., LTD."
    ws["J9"] = "2026-10-22"
    ws["J10"] = "SHSJ2610051A"
    ws["B17"] = "By Sea"
    ws["A18"] = "Port of loading: NINGBO, CHINA"
    ws["A36"] = "REF:26LF0845-1"
    ws["D24"] = "JC"
    ws["G24"] = 6000
    ws["D25"] = "JC"
    ws["G25"] = 664
    ws["D26"] = "POLY"
    ws["G26"] = 123
    ws["A38"] = "CONTAINER NO / SEAL NO:"
    ws["A39"] = "FBLU0186990 / SEAL123"
    ws["A40"] = "TEMU9126356 / SEAL456"
    ws["A42"] = "BY AIR"
    pkl = wb.create_sheet("PKL1")
    pkl["A24"] = 307
    pkl["B24"] = "PACKS"
    pkl["K24"] = 6900
    return wb


def main() -> None:
    fails: list[str] = []

    def check(cond: bool, label: str) -> None:
        print(("[OK] " if cond else "[X] ") + label)
        if not cond:
            fails.append(label)

    check(parse_date("2026-10-22") == date(2026, 10, 22), "解析 YYYY-MM-DD")
    check(parse_date("2026/10/22") == date(2026, 10, 22), "解析 YYYY/MM/DD")
    check(parse_date("2026-10-22 00:00:00") == date(2026, 10, 22), "解析带时间")
    check(parse_date("") is None, "空日期")

    check(mode_from("空运") == "AIR", "运输方式 空运→AIR")
    check(mode_from("海运") == "SEA", "运输方式 海运→SEA")
    check(mode_from("", "by Truck") == "TRUCK", "运输方式 Truck→TRUCK")

    check(
        _containers("", ["SZLU9863221/KAL0586", "PART OF 1X45HQ"])
        == ["SZLU9863221/KAL0586", "PART OF 1X45HQ"],
        "柜号原样逐格展示（不抽取不改写）",
    )
    check(
        _containers("A1\n\nB2", []) == ["A1", "", "B2"],
        "柜号按行对应 1-4 格，中间空行保留位置",
    )

    etd = date(2026, 10, 20)
    check(eta_for("AIR", etd) == date(2026, 10, 21), "ETA 空运 +1")
    check(eta_for("SEA", etd) == date(2026, 10, 28), "ETA 海运 +8")
    check(eta_for("TRUCK", etd) == date(2026, 10, 28), "ETA 汽运 +8")

    wb_txt = openpyxl.Workbook()
    ws_txt = wb_txt.active
    ws_txt["A1"] = 850
    ws_txt["B1"] = "PACKS"
    check(_read_packages(ws_txt) == 850, "取 PACKS 左侧单元格数字")
    ws_txt["C3"] = "307 PACKS"
    check(_read_packages(ws_txt) == 1157, "同格文本 307 PACKS 也取 packs 前数字")
    check(_read_packages(wb_txt.create_sheet("blank")) is None, "PKL 无 PACKS 时留空")
    wb_txt.close()

    row = build_row(
        [_source()],
        {
            "date": "2026-10-22",
            "bill_no": "SHSJ2610051A",
            "place_receipt": "Port of loading: Ningbo, China",
            "container_info": "FBLU0186990 / SEAL123",
        },
        company="Vanco",
        transport="海运",
    )
    check(row.custom_no == "SHSJ2610051A_CT", "Vanco 发票号后缀 _CT")
    check(row.blading == "SHSJ2610051A", "提单号取 bill_no")
    check(row.fty_inv == "26LF1025-1", "Fty Inv 取 PO No.")
    check(row.shipper == "LEFEI", "Shipper 取供应商代码")
    check(row.est_vco == "Vanco", "EST/VCO 取我司工厂")
    check(row.package == 307, "Package 取件数")
    check(row.containers == ["FBLU0186990 / SEAL123"], "柜号原样逐格展示")
    check(row.pol == "Ningbo, China", "POL 去掉标签前缀")
    check(row.etd == date(2026, 10, 21), "ETD 取发票日期")
    check(row.eta == date(2026, 10, 29), "ETA = ETD + 8")
    check(row.vessel == "SITC RUNDE V.2622S", "VESSEL 预填航次")
    check(row.carrier == "SITC", "Carrier 预填货代")
    check(month_of(row.date) == "2026-10", "月份归属")

    row_eastex = build_row(
        [_source(delivery_to="EASTEX INDUSTRIAL CO., LTD", supplier_code="SHUIXIN")],
        {"bill_no": "XMNCT2610214A"},
        company="Eastex",
        transport="陆运",
    )
    check(row_eastex.custom_no == "XMNCT2610214A_BT", "Eastex 发票号后缀 _BT")
    check(row_eastex.mode == "TRUCK", "陆运 Mode=TRUCK")

    merged = build_row(
        [
            _source(po_no="26LF1025-1", rolls=100, net_weight=1000, gross_weight=1100,
                    container_number="TEMU9126356"),
            _source(po_no="26LF1026", rolls=50, net_weight=500, gross_weight=550,
                    container_number="FBLU0186990"),
        ],
        {"bill_no": "2609007"},
        company="Vanco",
        transport="海运",
    )
    check(merged.fty_inv == "26LF1025-1, 26LF1026", "合并导出的 Fty Inv 用逗号连接")
    check(merged.package == 150, "合并件数求和")
    check(merged.nw_kg == 1500 and merged.gw_kg == 1650, "合并重量求和")
    check(merged.containers == ["TEMU9126356", "FBLU0186990"], "多柜号写入 1-2")
    check(merged.custom_no == "2609007_CT", "合并导出单号")

    res_wb = _result_workbook()
    try:
        res_row = build_row_from_result(
            res_wb, company="Vanco", transport="海运", customs_weight=6664
        )
        check(res_row.shipper == "ZHEJIANG LEFEI TEXTILE CO.,LTD.", "清关文件 Shipper 取 INV1!A1")
        check(res_row.est_vco == "Vanco", "清关文件 EST/VCO 取工厂")
        check(res_row.custom_no == "SHSJ2610051A_CT", "清关文件 Custom # 后缀")
        check(res_row.blading == "SHSJ2610051A", "清关文件 Blading 取发票号")
        check(res_row.fty_inv == "26LF0845-1", "清关文件 Fty Inv 取 REF")
        check(res_row.pol == "NINGBO, CHINA", "清关文件 POL 取 A18")
        check(
            res_row.containers
            == ["FBLU0186990 / SEAL123", "TEMU9126356 / SEAL456", "", "BY AIR"],
            "清关文件柜号逐格对应 Container 1-4（空格保留位置）",
        )
        check(res_row.nw_kg == 6664, "清关文件净重=报关净重")
        check(res_row.gw_kg == 6900, "清关文件毛重取 PKL1!K24")
        check(res_row.description == "JC / POLY", "清关文件分类去重")
        check(res_row.mode == "SEA", "清关文件运输方式")
        check(res_row.package == 307, "清关文件 Package 取 PKL PACKS 前的数字")
        check(res_row.etd is None and res_row.eta is None, "ETD/ETA 留空")
        check(res_row.vessel == "" and res_row.carrier == "", "VESSEL/Carrier 留空")
        check(res_row.container_type == "" and res_row.cbm is None, "柜型/CBM 留空待人工")
        check(
            res_row.machine_cabinet == "" and res_row.aux_cabinet == "",
            "机器柜/辅料柜 留空（仅 Shipment List 页手填）",
        )
    finally:
        res_wb.close()

    check(len(row.cells()) == 24, "cells() 返回 24 列")

    row.container_type = "40HQ"
    row.cbm = 45.5
    row.machine_cabinet = "MSKU1234567"
    row.aux_cabinet = "TGHU7654321"
    out = Path(tempfile.gettempdir()) / "shipment_list_test.xlsx"
    write_workbook([row, row_eastex], out)
    wb = openpyxl.load_workbook(out)
    try:
        ws = wb["Sheet1"]
        check(ws["A1"].value == "Date", "模板表头保留")
        check(ws["A2"].value.date() == date(2026, 10, 22), "第 2 行写入 Date")
        check(ws["E2"].value == "SHSJ2610051A_CT", "第 2 行写入 Custom #")
        check(ws["B3"].value == "26LF1025-1", "第 3 行写入 Fty Inv")
        check(ws["A6"].value is None, "模板说明行已清空，避免与数据冲突")
        check("mmm" in ws["A2"].number_format, "日期格式沿用模板")
        check(ws["A2"].border.left.style == "thin", "边框沿用模板")
        check(ws["U1"].value == "柜型" and ws["V1"].value == "CBM"
              and ws["W1"].value == "Machine Cabinet" and ws["X1"].value == "Aux Cabinet"
              and ws["Y1"].value == "ID",
              "新列表头 柜型/CBM/Machine Cabinet/Aux Cabinet/ID")
        check(ws["U2"].value == "40HQ" and ws["V2"].value == 45.5,
              "第 2 行写入柜型/CBM")
        check(ws["W2"].value == "MSKU1234567" and ws["X2"].value == "TGHU7654321",
              "第 2 行写入机器柜/辅料柜")
    finally:
        wb.close()
    out.unlink(missing_ok=True)

    # --- 装柜列表 ref_id → 发票明细反查映射 ---
    from sqlalchemy import create_engine as _create_engine
    from sqlalchemy.orm import sessionmaker as _sessionmaker

    from webapp.models import Base as _Base, ShipmentListEntry as _SLE
    from webapp.service import shipment_ref_map as _shipment_ref_map

    _engine = _create_engine("sqlite://")
    _Base.metadata.create_all(_engine)
    _Session = _sessionmaker(bind=_engine)
    _db = _Session()
    _db.add_all([
        _SLE(custom_no="A_CT", ref_id="ID-100", invoice_ids="[1,2]"),
        _SLE(custom_no="B_BT", ref_id="ID-200", invoice_ids="[3]"),
        _SLE(custom_no="C_CT", ref_id="", invoice_ids="[4]"),
        _SLE(custom_no="D_CT", ref_id="ID-NEW", invoice_ids="[1]"),
    ])
    _db.commit()
    _m = _shipment_ref_map(_db)
    check(_m.get(1) == "ID-NEW", "ref_id 反查：同发票取最新一条")
    check(_m.get(2) == "ID-100", "ref_id 反查：按发票 id 映射")
    check(_m.get(3) == "ID-200", "ref_id 反查：另一票映射")
    check(_m.get(4) is None, "ref_id 反查：未填 ID 不映射")
    _db.close()
    _engine.dispose()

    # --- record_shipment_list 的 ref_id 落库语义（弹窗手填 → 仅存数据库） ---
    import webapp.service as _service
    from webapp.models import Invoice as _Invoice
    from webapp.service import record_shipment_list as _record

    _orig_rebuild = _service.rebuild_shipment_month
    _service.rebuild_shipment_month = (
        lambda db, month: Path(tempfile.gettempdir()) / f"shipment_list_{month}.xlsx"
    )
    try:
        _engine2 = _create_engine("sqlite://")
        _Base.metadata.create_all(_engine2)
        _db2 = _sessionmaker(bind=_engine2)()
        _inv = _Invoice(invoice_no="T-1", source_filename="t.xlsm", source_file="t.xlsm")
        _db2.add(_inv)
        _db2.commit()
        _row = ShipmentRow(
            date=date(2026, 10, 22),
            fty_inv="PO-1",
            shipper="LEFEI",
            est_vco="Vanco",
            custom_no="T-1_CT",
        )
        _record(_db2, None, [_inv], _row, ref_id="ID-777")
        _e = _db2.query(_SLE).one()
        check(_e.ref_id == "ID-777", "导出弹窗 ref_id 写入装柜记录")
        _record(_db2, None, [_inv], _row, ref_id="")
        _db2.refresh(_e)
        check(_e.ref_id == "ID-777", "ref_id 留空保留原值（重导出不抹掉）")
        _record(_db2, None, [_inv], _row, ref_id="ID-888")
        _db2.refresh(_e)
        check(_e.ref_id == "ID-888", "ref_id 再次填写可覆盖")
        check(_db2.query(_SLE).count() == 1, "同一 source_key 重导出只更新不新增")
        check("ref_id" not in _row.as_dict(), "as_dict 不含 ref_id（Excel 行无此列）")
        # 机器柜/辅料柜仅 Shipment List 页手填：导出行恒为空，重导出不得覆盖
        _e.machine_cabinet = "MSKU1234567"
        _e.aux_cabinet = "TGHU7654321"
        _db2.commit()
        _record(_db2, None, [_inv], _row, ref_id="")
        _db2.refresh(_e)
        check(
            _e.machine_cabinet == "MSKU1234567" and _e.aux_cabinet == "TGHU7654321",
            "机器柜/辅料柜 重导出保留页面手填值",
        )
        check(
            "machine_cabinet" in _row.as_dict() and "aux_cabinet" in _row.as_dict(),
            "as_dict 含机器柜/辅料柜",
        )
    finally:
        _service.rebuild_shipment_month = _orig_rebuild
        _db2.close()
        _engine2.dispose()

    if fails:
        print(f"\n失败 {len(fails)} 项")
        raise SystemExit(1)
    print("\n全部通过")


if __name__ == "__main__":
    main()
