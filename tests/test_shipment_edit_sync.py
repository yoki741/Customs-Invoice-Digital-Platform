# -*- coding: utf-8 -*-
"""端到端：Shipment List 行内编辑（手动字段弹窗）→ 同步 .xlsm / 修改记录。

覆盖当前权限分割下的手工验证点：
1. 非手工行行内编辑仅手动字段（Package/ETD/ETA/VESSEL/Carrier/柜型/CBM/ID）；
   其中 Package 映射到 .xlsm PKL1 A24 → 改写文件 + file_history 留底 +
   修改记录时间线出现「编辑保存」含 old→new；越权字段（Date/柜号/NW/GW/
   Mode/POL 等）服务端保留原值、不碰文件、不写修改记录；
2. 改 ETD/ETA/VESSEL 等平台-only 字段 → 不碰文件、不插修改记录、月度 Excel 照常重建；
3. 改 ID → 明细页 ID 实时映射生效；
4. 平台改写后的文件回传校验不误报（比对带 sha256 的 save 行）；
5. 批量下载 zip 含 .xlsm。
"""

from __future__ import annotations

import io
import json
import re
import sys
import warnings
import zipfile
from pathlib import Path

warnings.simplefilter("ignore")
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from fastapi.testclient import TestClient  # noqa: E402

from webapp.main import app  # noqa: E402

SAMPLE = ROOT / "柬埔寨清关例子" / "VCO_202609_FB675_D202609007.xlsm"


def _csrf_from(html: str) -> str:
    m = re.search(r'name="csrf"\s+value="([^"]+)"', html)
    assert m, "未找到 CSRF token"
    return m.group(1)


def _xlsm_cells(path: Path) -> dict:
    """读取 .xlsm 关键单元格（柜号区/Date/Mode/PKL1/G 列合计）。"""
    import openpyxl

    wb = openpyxl.load_workbook(str(path), keep_vba=False)
    ws = wb["INV1"]
    container_row = None
    for r in range(24, 60):
        a_val = ws[f"A{r}"].value
        if a_val and "CONTAINER" in str(a_val).upper():
            container_row = r
            break
    cells = {
        "J9": ws["J9"].value,
        "B17": ws["B17"].value,
        "A18": ws["A18"].value,
        "containers": [
            ws[f"A{container_row + 1 + i}"].value if container_row else None
            for i in range(4)
        ],
    }
    if "PKL1" in wb.sheetnames:
        cells["PKL1_A24"] = wb["PKL1"]["A24"].value
        cells["PKL1_K24"] = wb["PKL1"]["K24"].value
    # G 列报关重量合计（面料行 = 24..Total 行-1，与改写分摊口径一致）
    total_row = 37
    for r in range(24, 60):
        c_val = ws[f"C{r}"].value
        if c_val and "Total" in str(c_val):
            total_row = r
            break
    g_sum = 0.0
    for r in range(24, total_row):
        val = ws[f"G{r}"].value
        if isinstance(val, (int, float)) and not isinstance(val, bool):
            g_sum += float(val)
    cells["G_sum"] = g_sum
    wb.close()
    return cells


def main() -> None:
    from webapp.models import (
        ExportFileVersion,
        ExportJob,
        Invoice,
        SessionLocal,
        ShipmentListEntry,
    )

    client = TestClient(app, base_url="http://testserver", follow_redirects=False)
    r = client.post(
        "/login",
        data={"username": "admin", "password": "admin123", "next": ""},
        follow_redirects=False,
    )
    assert r.status_code == 303, r.status_code

    # ---- 准备：导入样例 + 导出生成报关发票与装柜行 ----
    csrf = _csrf_from(client.get("/invoices").text)
    with SAMPLE.open("rb") as fh:
        r = client.post(
            "/admin/upload",
            data={"csrf": csrf},
            files=[("files", (SAMPLE.name, fh, "application/octet-stream"))],
        )
    if "Duplicate Invoice No." in r.text:
        csrf = _csrf_from(r.text)
        r = client.post(
            "/admin/upload",
            data={
                "csrf": csrf,
                "confirm": "1",
                "pending_json": json.dumps(
                    [{
                        "file": SAMPLE.name,
                        "path": str(SAMPLE),
                        "invoice_no": "D202609007",
                        "overwrite": True,
                    }]
                ),
            },
        )
    assert r.status_code == 200, r.text[:300]

    db = SessionLocal()
    inv = db.query(Invoice).filter(Invoice.invoice_no == "D202609007").first()
    assert inv is not None, "样例发票应存在"
    inv_id = inv.id

    payload = {
        "form": {
            "date": "2026-09-10",
            "bill_no": "2609007",
            "supplier": "VANCO",
            "ship_mode": "by Truck",
            "place_receipt": "SHAOXING,CHINA",
            "port_loading": "HUU NGHI BORDER, VIET NAM",
            "container_info": "SZLU9863221/KAL0586\nPART OF 1X45HQ",
            "flight_no": "",
        }
    }
    r = client.post("/api/invoices/export", params={"ids": str(inv_id)}, json=payload)
    assert r.status_code == 200 and r.json()["ok"], r.text[:500]
    exported = r.json()
    folder, filename = exported["folder"], exported["filename"]
    out_path = Path(exported["path"])
    assert out_path.exists()

    job = db.query(ExportJob).order_by(ExportJob.id.desc()).first()
    entry = (
        db.query(ShipmentListEntry)
        .filter(ShipmentListEntry.job_id == job.id)
        .order_by(ShipmentListEntry.id.desc())
        .first()
    )
    assert entry is not None, "导出后应有装柜记录"
    entry_id = entry.id

    # 越权字段的导出原值（行内编辑对非手工行不得改动，保存后应原样保留）
    keep = {
        f: getattr(entry, f)
        for f in ("date", "container_1", "container_2", "nw_kg", "gw_kg", "mode", "pol")
    }
    cells_before = _xlsm_cells(out_path)

    hist_dir = ROOT / "data" / "file_history" / folder
    hist_before = len(list(hist_dir.glob("*.xlsm"))) if hist_dir.exists() else 0

    # ---- 1. 行内编辑：仅手动字段（Package/ETD/ETA/VESSEL/ID）可写 ----
    form = {
        "date": "2026-09-12",          # 越权字段，应被忽略
        "fty_inv": entry.fty_inv or "",
        "shipper": entry.shipper or "",
        "est_vco": entry.est_vco or "",
        "custom_no": entry.custom_no or "",
        "package": "50",               # 手动字段且映射 .xlsm PKL1 A24
        "blading": entry.blading or "",
        "container_1": "NEWCONT1/SEAL1",  # 越权字段，应被忽略
        "container_2": "",
        "container_3": "",
        "container_4": "",
        "nw_kg": "1234",               # 越权字段，应被忽略
        "gw_kg": "2000",               # 越权字段，应被忽略
        "description": entry.description or "",
        "mode": "SEA",                 # 越权字段，应被忽略
        "etd": "2026-09-20",
        "eta": "",
        "pol": "NINGBO, CHINA",        # 越权字段，应被忽略
        "vessel": "NEW VESSEL",
        "carrier": entry.carrier or "",
        "container_type": entry.container_type or "",
        "cbm": "",
        "ref_id": "TEST-001",
    }
    r = client.post(f"/api/shipments/{entry_id}/update", data=form)
    assert r.status_code == 200, r.text[:500]
    result = r.json()
    assert result["ok"], result
    assert result["xlsm_synced"] is True, result  # Package 改写 .xlsm
    print("[OK] 行内编辑保存:", result["message"])

    db.expire_all()
    entry = db.get(ShipmentListEntry, entry_id)
    assert entry.package == 50, entry.package
    assert entry.vessel == "NEW VESSEL" and entry.ref_id == "TEST-001"
    assert entry.eta == "2026-09-28", entry.eta  # TRUCK: ETD+8 自动推算
    # 越权字段保留导出原值
    for f, val in keep.items():
        assert getattr(entry, f) == val, f"{f} 不得被行内编辑改动"
    assert entry.date == "2026-09-10", entry.date
    print("[OK] 手动字段生效 + ETA 自动推算；Date/柜号/NW/GW/Mode/POL 保留原值")

    cells = _xlsm_cells(out_path)
    assert cells["PKL1_A24"] == 50, cells["PKL1_A24"]  # Package 同步
    assert cells["J9"] == cells_before["J9"], cells["J9"]  # Date 未同步
    assert cells["B17"] == cells_before["B17"], cells["B17"]  # Mode 未同步
    assert cells["containers"] == cells_before["containers"], cells["containers"]
    assert cells["PKL1_K24"] == cells_before["PKL1_K24"], cells["PKL1_K24"]  # GW 未同步
    assert cells["G_sum"] == cells_before["G_sum"], cells["G_sum"]  # NW 未同步
    print("[OK] .xlsm 仅改写 PKL1 A24（Package），其余单元格不动")

    hist_after = len(list(hist_dir.glob("*.xlsm"))) if hist_dir.exists() else 0
    assert hist_after == hist_before + 1, (hist_before, hist_after)
    print("[OK] 旧文件已留底 file_history +1")

    save_row = (
        db.query(ExportFileVersion)
        .filter(ExportFileVersion.kind == "save", ExportFileVersion.folder == folder)
        .order_by(ExportFileVersion.id.desc())
        .first()
    )
    assert save_row is not None, "应有改写留痕 save 行"
    assert save_row.sha256 and save_row.backup_path, save_row.__dict__
    assert Path(save_row.backup_path).exists()
    snapshot = json.loads(save_row.form_values or "{}")
    changed_keys = {c["key"] for c in snapshot.get("_changes") or []}
    assert "package" in changed_keys, changed_keys
    assert "date" not in changed_keys and "container_1" not in changed_keys, changed_keys
    print("[OK] 修改记录 save 行：sha256 + 留底 + 仅 Package 变更快照")

    r = client.get("/api/files/history", params={"folder": folder, "name": filename})
    assert r.status_code == 200 and r.json()["ok"], r.text[:300]
    rows = r.json()["rows"]
    last = rows[-1]
    assert last["kind"] == "save", last
    assert last["backup"], "save 行带留底下载"
    changed_labels = {c["label"] for c in last["changes"]}
    assert "Package" in changed_labels, changed_labels
    assert "Date" not in changed_labels and "Container 1" not in changed_labels, changed_labels
    assert any(c["old"] and c["new"] for c in last["changes"]), last["changes"]
    print("[OK] 历史时间线：编辑保存 + Package old→new:", changed_labels)

    # 发票明细 ID 同步（柜号已不再可由行内编辑改动，manual_values 保持原值）
    r = client.get("/details", params={"invoice_id": inv_id})
    assert r.status_code == 200
    assert "TEST-001" in r.text, "明细页展示新 ID"
    print("[OK] 明细页 ID 实时映射生效")

    # ---- 2. 平台-only 字段（ETD/ETA/VESSEL）：不碰文件、不插记录 ----
    saves_before = (
        db.query(ExportFileVersion)
        .filter(ExportFileVersion.kind == "save", ExportFileVersion.folder == folder)
        .count()
    )
    form2 = dict(form)
    form2["etd"] = "2026-09-25"
    form2["eta"] = "2026-10-01"
    form2["vessel"] = "OTHER VESSEL"
    r = client.post(f"/api/shipments/{entry_id}/update", data=form2)
    assert r.status_code == 200 and r.json()["ok"], r.text[:300]
    assert r.json()["xlsm_synced"] is False, r.json()
    saves_after = (
        db.query(ExportFileVersion)
        .filter(ExportFileVersion.kind == "save", ExportFileVersion.folder == folder)
        .count()
    )
    assert saves_after == saves_before, (saves_before, saves_after)
    assert _xlsm_cells(out_path)["PKL1_A24"] == 50
    db.expire_all()
    entry = db.get(ShipmentListEntry, entry_id)
    assert entry.etd == "2026-09-25" and entry.eta == "2026-10-01"
    print("[OK] ETD/ETA/VESSEL 平台-only：文件不动、无新记录")

    # ---- 3. 回传校验：平台改写后的文件与 save 行一致（不误报） ----
    with out_path.open("rb") as fh:
        r = client.post(
            "/api/files/verify",
            files=[("file", (filename, fh, "application/octet-stream"))],
        )
    assert r.status_code == 200, r.text[:300]
    data = r.json()
    assert data["matched"] is True, data
    assert "edit-save version" in data["message"], data["message"]
    print("[OK] 校验：改写后的文件与平台 save 记录一致")

    # ---- 4. 批量下载 zip ----
    r = client.get("/api/shipments/download-files", params={"ids": str(entry_id)})
    assert r.status_code == 200, r.text[:300]
    zf = zipfile.ZipFile(io.BytesIO(r.content))
    names = zf.namelist()
    assert filename in names, names
    assert "missing.txt" not in names
    print("[OK] 批量下载 zip 含", names)

    # ---- 5. 月度 Excel 照常重建 ----
    from core.shipment_list import month_path

    month_file = month_path(entry.month)
    assert month_file.exists(), month_file
    print("[OK] 月度 Excel 存在:", month_file.name)

    db.close()
    print("\nShipment List 行内编辑同步测试全部通过")


if __name__ == "__main__":
    main()
