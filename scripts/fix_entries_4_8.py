# -*- coding: utf-8 -*-
"""修复清空测试数据后残留的污染：id=4(D202609007) 与 id=8(26LF0846)。

污染来源（审计日志已确认）：
- id=4 被测试用 bill_no=2609007 重新导出，blading/custom/container 全被写坏，
  job_id 指向已删的测试任务 161；真实值取自真实导出 job 88（2026-10-22 / TRUCK /
  HUU NGHI BORDER / bill_no=D202609007 / 空柜号）。
- id=8 被测试导出成 26LF0846-CT（job 160），且手工字段被填入 245/358/27457/28522
  等假数字；真实值取自 job 159（SEA / SHANGHAI / 26LF0846 / 空柜号）。
- Invoice.manual_values 与 .xlsm 文件柜号单元格也残留 TESTSAVE0001 / 2609007。
"""
from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402

from webapp.models import (  # noqa: E402
    ExportFileVersion,
    Invoice,
    SessionLocal,
    ShipmentListEntry,
)
from webapp.service import rebuild_shipment_month  # noqa: E402


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    h.update(path.read_bytes())
    return h.hexdigest()


def main() -> None:
    db = SessionLocal()

    # ---- id=4 (Vanco:D202609007) 恢复真实值 ----
    e4 = db.get(ShipmentListEntry, 4)
    assert e4 is not None and e4.source_key == "Vanco:D202609007"
    e4.job_id = 88
    e4.month = "2026-10"
    e4.date = "2026-10-22"
    e4.blading = "D202609007"
    e4.custom_no = "D202609007_CT"
    e4.container_1 = ""
    e4.container_2 = ""
    print("id=4 ->", e4.date, e4.mode, e4.pol, e4.blading, e4.custom_no, repr(e4.container_1))

    # ---- id=8 (Eastex:26LF0846) 恢复真实值 + 清假数字手工字段 ----
    e8 = db.get(ShipmentListEntry, 8)
    assert e8 is not None and e8.source_key == "Eastex:26LF0846"
    e8.job_id = 159
    e8.blading = "26LF0846"
    e8.custom_no = "26LF0846_BT"
    e8.vessel = ""
    e8.carrier = ""
    e8.container_type = ""
    e8.cbm = None
    e8.etd = ""
    e8.eta = ""
    print("id=8 ->", e8.blading, e8.custom_no, repr(e8.vessel), repr(e8.carrier), repr(e8.container_type))

    # ---- id=3 清假数字手工字段（vessel/carrier/柜型/cbm，格式明显不符） ----
    e3 = db.get(ShipmentListEntry, 3)
    if e3 is not None:
        e3.vessel = ""
        e3.carrier = ""
        e3.container_type = ""
        e3.cbm = None
        print("id=3 清空假数字手工字段")

    # ---- Invoice.manual_values 同步 ----
    inv9 = db.get(Invoice, 9)
    if inv9 is not None:
        mv = json.loads(inv9.manual_values or "{}")
        mv["bill_no"] = "D202609007"
        mv["date"] = "2026-10-22"
        mv["container_info"] = ""
        inv9.manual_values = json.dumps(mv, ensure_ascii=False)

    inv2 = db.get(Invoice, 2)
    if inv2 is not None:
        mv = json.loads(inv2.manual_values or "{}")
        mv["bill_no"] = "26LF0846"
        mv["container_info"] = ""
        inv2.manual_values = json.dumps(mv, ensure_ascii=False)
    print("Invoice 9/2 manual_values 已同步")

    db.commit()

    # ---- 清 .xlsm 文件柜号单元格的 TESTSAVE0001 ----
    fpath = ROOT / "data" / "outputs" / "26LF0846" / "EASTEX INV#26LF0846-BT(SEA).xlsm"
    if fpath.exists():
        wb = openpyxl.load_workbook(str(fpath), keep_vba=False)
        ws = wb["INV1"]
        crow = None
        for r in range(24, 60):
            a = ws[f"A{r}"].value
            if a and "CONTAINER" in str(a).upper():
                crow = r
                break
        if crow is not None and ws[f"A{crow + 1}"].value == "TESTSAVE0001":
            ws[f"A{crow + 1}"] = None
            wb.save(str(fpath))
            new_sha = _sha256(fpath)
            # 更新该文件最新 export 记录的 sha256，避免校验误报
            ver = (
                db.query(ExportFileVersion)
                .filter(
                    ExportFileVersion.folder == "26LF0846",
                    ExportFileVersion.filename == "EASTEX INV#26LF0846-BT(SEA).xlsm",
                    ExportFileVersion.kind == "export",
                )
                .order_by(ExportFileVersion.id.desc())
                .first()
            )
            if ver is not None:
                ver.sha256 = new_sha
                db.commit()
            print("已清 26LF0846 .xlsm 柜号单元格 TESTSAVE0001 并更新 sha256")
        wb.close()

    # ---- 重建月度 Excel ----
    for month in ("2026-08", "2026-09", "2026-10"):
        try:
            rebuild_shipment_month(db, month)
            print(f"重建月度 Excel: {month}")
        except Exception as exc:  # noqa: BLE001
            print(f"重建 {month} 失败: {exc}")

    db.close()
    print("修复完成")


if __name__ == "__main__":
    main()
